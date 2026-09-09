from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "NeuroScan_AI_Academic_Project_Report.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.9); sec.right_margin = Inches(0.9)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"; styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.line_spacing = 1.35
styles["Normal"].paragraph_format.space_after = Pt(7)
for s in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    styles[s].font.name = "Times New Roman"; styles[s]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    styles[s].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(20); styles["Title"].font.bold = True
styles["Heading 1"].font.size = Pt(15); styles["Heading 1"].font.bold = True
styles["Heading 2"].font.size = Pt(12); styles["Heading 2"].font.bold = True

def p(text="", style=None, bold_prefix=None):
    para = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = para.add_run(bold_prefix); r.bold = True
        para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style not in ("Title",) else WD_ALIGN_PARAGRAPH.CENTER
    return para

def heading(text, level=1):
    para = doc.add_paragraph(text, style=f"Heading {level}")
    para.paragraph_format.space_before = Pt(14); para.paragraph_format.space_after = Pt(6)
    return para

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=c._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"1F4E79"); tcPr.append(shd)
        for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t

def bullet(text): doc.add_paragraph(text, style="List Bullet")

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("NeuroScan AI Academic Project Report | Research Use Only")
footer.runs[0].font.size = Pt(8)

# Title page
for _ in range(5): doc.add_paragraph()
title=p("NeuroScan AI", "Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub=p("Brain Tumor Segmentation and Visualization Using Multi Modal MRI and 3D U Net", "Heading 1"); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(5): doc.add_paragraph()
for line in ["Academic Project Report", "", "Submitted by", "Hamza Mustafa", "", "Degree Program: ______________________________", "Department: _________________________________", "Institution: __________________________________", "Supervisor: __________________________________", "Submission Date: ______________________________"]:
    q=doc.add_paragraph(line); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if q.runs: q.runs[0].font.size=Pt(12)
doc.add_page_break()

heading("Abstract")
p("NeuroScan AI is a local full stack research application for multi modal brain MRI tumor segmentation. The system accepts four conventional MRI modalities, performs preprocessing, applies a 3D U Net segmentation model when a trained checkpoint is available, and presents the resulting tumor sub regions in an interactive NIfTI viewer. The project also provides a safe demonstration mode, training metrics, PDF reporting, and local case management. The implementation combines a Next.js frontend with a FastAPI inference service and a PyTorch based model workflow aligned with the BraTS 2021 task. The current report documents the problem, system architecture, methods, implementation decisions, validation approach, limitations, and future work. The software is intended for research, learning, and visualization only; it is not a clinical decision support system or medical device.")
heading("Keywords", 2); p("brain tumor segmentation; magnetic resonance imaging; deep learning; 3D U Net; BraTS; FastAPI; Next.js; NIfTI; DICOM")
heading("Declaration", 2); p("This report describes an academic software project. Any model outputs, demonstration masks, measurements, or generated reports must not be used for diagnosis, treatment planning, or independent clinical decision making.")
doc.add_page_break()

heading("1 Introduction")
p("Brain tumor analysis from magnetic resonance imaging requires the interpretation of volumetric information across complementary image sequences. Manual outlining of tumor regions is time consuming and can vary between observers. Automated segmentation can support research workflows by producing reproducible region masks and quantitative summaries. The BraTS benchmark established a widely used setting for comparing computational methods on multi parametric MRI and tumor sub region segmentation [1].")
heading("1.1 Problem Statement", 2)
p("A practical research workflow needs more than a trained neural network. It needs controlled ingestion of image data, modality validation, preprocessing consistent with training, inference, a transparent display of model output, and a usable way to inspect and export results. Many student projects stop at a notebook prediction; NeuroScan AI addresses the wider engineering workflow while retaining local execution.")
heading("1.2 Objectives", 2)
for x in ["Develop a local web application for uploading and managing multi modal MRI studies.", "Implement a preprocessing and 3D U Net inference path compatible with the project training script.", "Visualize MRI volumes and segmentation masks across orthogonal views.", "Provide quantitative voxel and volume summaries for tumor related regions.", "Harden local uploads against unsafe archive paths and excessive archive expansion.", "Document clear non clinical and demonstration mode boundaries."] : bullet(x)
heading("1.3 Scope", 2)
p("The system supports research input in NIfTI format and ZIP based DICOM studies whose series descriptions or folder names clearly identify T1, T1CE, T2, and FLAIR. It does not provide user accounts, clinical interoperability certification, hospital deployment controls, or prospective clinical validation.")

heading("2 Background and Related Work")
p("U Net introduced an encoder decoder architecture with skip connections for biomedical image segmentation [2]. Its volumetric extension, 3D U Net, replaces two dimensional operations with three dimensional operations to learn dense segmentation from volumetric input [3]. This project adopts a compact 3D U Net variant to balance memory use with end to end volumetric processing.")
p("BraTS 2021 evaluates algorithms on pre operative multi parametric MRI and includes the segmentation of histologically distinct tumor sub regions [1]. NeuroScan AI follows the conventional input modalities of T1, contrast enhanced T1, T2, and FLAIR. The implementation represents necrotic or non enhancing tumor core, peritumoral edema, and enhancing tumor, while whole tumor and tumor core are calculated as derived composite measurements.")

heading("3 System Requirements and Design")
heading("3.1 Functional Requirements", 2)
table(["ID", "Requirement"], [["FR1","Accept a complete four modality NIfTI study or one ZIP study."],["FR2","Validate modality availability before segmentation."],["FR3","Run a trained 3D U Net checkpoint or explicitly labeled demonstration output."],["FR4","Display MRI volumes and segmentation overlays interactively."],["FR5","Calculate voxel counts and physical volumes for regions."],["FR6","Export a PDF summary and a segmentation mask."],["FR7","Accept a ZIP based DICOM study and convert identified series to temporary NIfTI."]], [0.55,6.35])
heading("3.2 Non Functional Requirements", 2)
table(["Area", "Design Response"], [["Privacy","Local processing; uploaded source DICOM files are temporary during conversion."],["Safety","Research use disclaimer and explicit synthetic demonstration labeling."],["Security","Case id validation; safe ZIP paths; file count, size, and compression ratio limits."],["Usability","Guided upload checklist; modality status indicators; responsive clinical research UI."],["Maintainability","Frontend and backend are separated; model architecture is mirrored in training and inference code."]], [1.3,5.6])

heading("4 Architecture")
p("The application has three main layers. The presentation layer is a Next.js application using TypeScript, Tailwind CSS, and NiiVue. It provides the dashboard, uploader, case detail workspace, metrics page, and local status feedback. The service layer is a FastAPI application that exposes REST endpoints for health, cases, prediction, metrics, image volumes, masks, and reports. The machine learning layer uses PyTorch and a compact 3D U Net model. Training is designed to run in a Kaggle environment, while local inference loads the exported model weights.")
table(["Component", "Technology", "Responsibility"], [["Frontend","Next.js, React, TypeScript","User interface, upload interaction, viewer controls, charts."],["Backend","FastAPI, Pydantic","API validation, storage, inference coordination, reports."],["Imaging","NiBabel, NiiVue, PyDicom","NIfTI reading, browser visualization, DICOM series conversion."],["ML","PyTorch","3D U Net inference, loss computation, training checkpointing."],["Reporting","ReportLab, Pillow","PDF report and representative slice images."]], [1.2,1.8,3.9])

heading("5 Methodology")
heading("5.1 Data Preparation", 2)
p("For each study, the system requires T1, T1CE, T2, and FLAIR. NIfTI files are normalized to a consistent .nii.gz storage convention. For a DICOM ZIP, the importer groups files by Series Instance UID, identifies modality hints from Series Description, Protocol Name, or folder name, stacks decoded slices, builds an affine approximation from DICOM orientation and spacing metadata, and writes temporary NIfTI files. Ambiguous or incomplete studies fail with an actionable error instead of silently guessing.")
heading("5.2 Preprocessing", 2)
p("Each modality is loaded as a float32 volume. Intensities are z score normalized over non zero brain voxels while the background remains zero. The volumes are center cropped or zero padded to 128 by 128 by 128, matching the training configuration. The stored transformation is inverted after inference so that predicted labels are returned in the original volume dimensions.")
heading("5.3 Model and Training", 2)
p("The model uses an encoder decoder 3D U Net with four input channels and three output channels. Convolution blocks use 3 by 3 by 3 convolutions, instance normalization, and LeakyReLU activations. Downsampling uses strided convolution and upsampling uses transposed convolution with skip connections. Training combines binary cross entropy with Dice loss. The model is trained with a batch size of one and a cosine annealing learning rate schedule, reflecting the memory requirements of 128 cubed volumetric input.")
heading("5.4 Postprocessing and Quantification", 2)
p("Model logits are passed through sigmoid activation. At each voxel, the implementation assigns the highest scoring class above a configured threshold to produce a BraTS compatible label map. Per label volumes use voxel spacing from the reference image. Whole tumor is computed as the union of all non background labels; tumor core is computed as necrotic core plus enhancing tumor.")

heading("6 Implementation")
heading("6.1 Upload and Local Safety", 2)
p("The upload endpoint streams data to disk rather than trusting browser supplied file size metadata. It applies an aggregate limit controlled by NEUROSCAN_MAX_UPLOAD_BYTES, with a default of 2 GB. ZIP files are rejected if they contain unsafe paths, too many members, an excessive decompressed size, or a suspicious compression ratio. Server generated case identifiers are validated before filesystem operations. These measures address common local archive ingestion risks, although they are not a replacement for enterprise security controls.")
heading("6.2 Inference Modes", 2)
p("When model.pt is available, the backend loads tensor weights using PyTorch weights only deserialization and executes trained inference. When no checkpoint is present, the system enters a clearly stated demonstration mode. In that mode it generates a deterministic synthetic nested mask only to exercise the upload, viewer, statistics, and report workflow. The interface labels this output as non clinical and not representative of the uploaded scan.")
heading("6.3 User Interface", 2)
p("The interface was redesigned as a light clinical research workspace. The landing screen communicates a three stage workflow: prepare, segment, and review. The upload page explains the required modalities and permitted formats. The case workspace offers modality selection, slice orientation selection, segmentation opacity, region toggles, report export, and mask download. Color labels are synchronized across training metrics, backend statistics, and viewer overlays.")

heading("7 Testing and Evaluation")
p("Software verification focused on build integrity and workflow level behavior. The frontend passed ESLint, TypeScript checking, and a production Next.js build. Python source compilation passed for the backend and training modules. The redesigned dashboard and upload workflow were inspected in a browser. The backend integration environment was not available during the final audit because FastAPI, NiBabel, and PyDicom were not installed in the workspace; therefore, real DICOM conversion should be exercised after dependency installation using a de identified test study.")
table(["Test Area", "Result", "Evidence"], [["Frontend lint","Pass","ESLint completed with no errors."],["Type checking","Pass","TypeScript completed with no errors."],["Production build","Pass","Next.js generated all application routes."],["Python syntax","Pass","Backend and training modules compiled."],["UI review","Pass","Dashboard and upload screens rendered and were manually inspected."],["Real DICOM conversion","Pending environment setup","Requires installed backend dependencies and de identified DICOM series."]], [1.6,1.3,4.0])
heading("7.1 Evaluation Metrics", 2)
p("The training script reports Dice score per tumor class, train loss, and validation loss. Dice score measures spatial overlap between prediction and ground truth: Dice equals two times the intersection divided by the sum of predicted and ground truth voxels. The application deliberately presents training time Dice as validation context, not as a per patient accuracy score.")

heading("8 Limitations and Ethical Considerations")
p("The project is a local academic prototype. It has not undergone clinical validation, prospective testing, calibration analysis, bias assessment, regulatory review, or cybersecurity certification. DICOM support is intentionally constrained to clearly labeled, conventional series in one ZIP archive. The affine construction is suitable for research visualization but should not be interpreted as a certified clinical image conversion pipeline. Clinical adoption would require authentication, access controls, encrypted storage, audit logging, robust DICOM interoperability, model governance, independent validation, and institutional approval.")

heading("9 Conclusion and Future Work")
p("NeuroScan AI demonstrates an end to end approach to brain tumor segmentation research: data ingestion, preprocessing, 3D U Net inference, visualization, quantification, and reporting are integrated into one local application. The project contributes a usable software layer around a volumetric segmentation model and explicitly separates demonstration behavior from trained model predictions. Future work should include reproducible dataset splits, cross validation, automated test fixtures for NIfTI and DICOM, model calibration and uncertainty estimation, richer quality assurance, DICOM de identification workflows, and a formal evaluation on held out BraTS data.")

heading("References")
refs=["[1] U. Baid et al., The RSNA ASNR MICCAI BraTS 2021 Benchmark on Brain Tumor Segmentation and Radiogenomic Classification, arXiv:2107.02314, 2021.","[2] O. Ronneberger, P. Fischer, and T. Brox, U Net Convolutional Networks for Biomedical Image Segmentation, MICCAI, pp. 234 241, 2015.","[3] O. Cicek, A. Abdulkadir, S. S. Lienkamp, T. Brox, and O. Ronneberger, 3D U Net Learning Dense Volumetric Segmentation from Sparse Annotation, MICCAI, pp. 424 432, 2016.","[4] NeuroScan AI source code and project documentation, local project repository, accessed September 2026.","[5] NiiVue documentation and source repository, https://github.com/niivue/niivue, accessed September 2026."]
for r in refs: p(r)

doc.add_page_break(); heading("Appendix A Project Structure")
table(["Directory", "Contents"], [["frontend","Next.js pages, UI components, NIfTI viewer, charts, client API."],["backend","FastAPI routes, case storage, preprocessing, inference, reports, DICOM importer."],["training","Kaggle ready 3D U Net training script and metrics export."],["docs","Design system and project documentation."],["backend/models_store","Trained model weights and metrics when available."]], [1.7,5.2])
heading("Appendix B Reproduction Steps", 2)
for x in ["Create a Python virtual environment in backend and install requirements.txt.", "Start the backend using uvicorn app.main:app --reload --port 8000.", "Install frontend dependencies and run npm run dev in frontend.", "Open http://localhost:3000 and upload a de identified four modality NIfTI case or valid DICOM ZIP.", "For trained inference, copy model.pt and metrics.json into backend/models_store and restart the backend."]: bullet(x)
heading("Appendix C Submission Checklist", 2)
for x in ["Replace institution, degree program, supervisor, and date on the title page.", "Add your university required cover page or declaration if applicable.", "Insert actual experiment results and screenshots from your trained model before claiming performance.", "Keep the research use disclaimer with any demonstration mode screenshots.", "Verify the required citation style against your department guidelines."]: bullet(x)

doc.core_properties.author = "Hamza Mustafa"
doc.core_properties.title = "NeuroScan AI Academic Project Report"
doc.save(OUT)
print(OUT)
